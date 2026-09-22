#!/usr/bin/env python3
"""Rebuild the manuscript without touching project code or frozen evidence."""
import argparse
import hashlib
import os
from pathlib import Path
import re
import shutil
import subprocess
import zipfile

ROOT = Path(__file__).resolve().parent
FIGURE2_SHA256 = 'e8cfa3b68930c9052c7fbe4190eb7841efb428ab1bfe587e03dbea9fac0d1f24'


def run(args, cwd, log, env=None):
    with log.open('w') as output:
        subprocess.run(args, cwd=cwd, stdout=output, stderr=subprocess.STDOUT,
                       env=env, check=True)


def compile_tex(source, build, bibliography=False):
    stem = Path(source).stem
    command = ['pdflatex', '-interaction=nonstopmode', '-halt-on-error',
               '-file-line-error', f'-output-directory={build}', source]
    run(command, ROOT, build / f'{stem}-pass1.txt')
    if bibliography:
        env = dict(os.environ, BIBINPUTS=str(ROOT) + os.pathsep)
        run(['bibtex', stem], build, build / f'{stem}-bibtex.txt', env)
        run(command, ROOT, build / f'{stem}-pass2.txt')
    run(command, ROOT, build / f'{stem}-final.txt')
    log = (build / f'{stem}.log').read_text()
    failures = re.findall(
        r'Overfull \\[hv]box[^\n]*|'
        r'LaTeX Warning: Reference [^\n]* undefined[^\n]*|'
        r'Package natbib Warning: Citation [^\n]* undefined[^\n]*|'
        r'(?:LaTeX|Package natbib) Warning: There were undefined (?:references|citations)[^\n]*|'
        r'[^\n]*multiply defined[^\n]*|Missing character[^\n]*',
        log)
    if failures:
        raise RuntimeError('Typesetting validation failed:\n' + '\n'.join(failures))


def export_word(build):
    # Pandoc does not parse tabularx or TikZ. Convert only the export input;
    # the authoritative ACM LaTeX source remains unchanged.
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    source = (ROOT / 'main.tex').read_text()
    abstract = re.search(r'\\begin\{abstract\}(.*?)\\end\{abstract\}', source, re.S)[1]
    source = source.replace(r'\maketitle', r'\maketitle' + '\n' +
                            r'\noindent swayam.singal@gmail.com' + '\n\n' +
                            r'\section*{Abstract}' + abstract)
    source = source.replace(r'\paragraph{', r'\paragraph*{')
    source = source.replace(r'\mathrm{Pred}_{\hard}', r'\mathrm{Pred}_{\mathrm{HARD}}')
    definitions = source[source.index(r'\newcommand{\method}'):source.index(r'\title{')]
    diagram = build / 'system_flow_export.tex'
    diagram.write_text(r'''\documentclass[border=4pt]{standalone}
\usepackage{amsmath,tikz,tabularx}
\usepackage{lmodern}
\usetikzlibrary{arrows.meta,positioning,calc}
\setlength{\textwidth}{7in}
''' + definitions + r'\begin{document}\input{figures/system_flow.tex}\end{document}')
    compile_tex(str(diagram), build)
    run(['pdftoppm', '-singlefile', '-r', '220', '-png',
         str(build / 'system_flow_export.pdf'), str(build / 'system_flow_export')],
        ROOT, build / 'system-flow-render.txt')
    source = source.replace(r'\input{figures/system_flow.tex}',
                            r'\includegraphics{' + str(build / 'system_flow_export.png') + '}')
    numbers = dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}',
                              (build / 'main.aux').read_text()))
    source = re.sub(r'\\ref\{([^}]+)\}', lambda m: numbers[m[1]], source)
    lines = source.splitlines()
    for index, line in enumerate(lines):
        if line.startswith(r'\begin{tabularx}'):
            end = next(i for i in range(index + 1, len(lines))
                       if lines[i].startswith(r'\end{tabularx}'))
            header = next(v for v in lines[index + 1:end] if '&' in v)
            columns = header.count('&') + 1
            lines[index] = r'\begin{tabular}{' + 'l' * columns + '}'
            lines[end] = r'\end{tabular}'
    source = '\n'.join(lines)
    source = re.sub(r'\\shortstack\{([^{}]*(?:\\FULL\{\}[^{}]*)?)\}',
                    lambda m: m[1].replace(r'\\', ' '), source)
    source = source.replace('figure*', 'figure').replace('table*', 'table')
    while r'\caption*{' in source:
        start = source.index(r'\caption*{')
        pos = start + len(r'\caption*{')
        depth, end = 1, pos
        while depth:
            if source[end] == '{':
                depth += 1
            elif source[end] == '}':
                depth -= 1
            end += 1
        source = source[:start] + r'\par ' + source[pos:end - 1] + r'\par' + source[end:]
    # Word/Pandoc export: leave display equations unnumbered during conversion.
    # Plain Word equation numbers are added beside editable OMML equations below.
    equation_count = len(re.findall(r'\\begin\{equation\}', source))

    # Pandoc/LibreOffice do not reliably render AMS aligned markers in OMML.
    # Flatten the two multi-line displays for the editable Word export only.
    criticality_aligned = r'''\begin{aligned}
\chi(C)&=\max_q r_q(C),\\
r_q(C)&=
\begin{cases}
 B_q(C)/\varepsilon_q, & \varepsilon_q>0,\\
 0, & \varepsilon_q=B_q(C)=0,\\
 +\infty, & \varepsilon_q=0<B_q(C).
\end{cases}
\end{aligned}'''
    criticality_flat = r'''\chi(C)=\max_q r_q(C),\qquad
r_q(C)=\begin{cases}
 B_q(C)/\varepsilon_q, & \varepsilon_q>0,\\
 0, & \varepsilon_q=B_q(C)=0,\\
 +\infty, & \varepsilon_q=0<B_q(C).
\end{cases}'''
    gaussian_aligned = r'''\begin{aligned}
\Delta_p
&=\left\lVert R_p(U\cup E_0)-R_p(U\cup E_1)\right\rVert_\infty,\\
\Delta_p
&\le C_{\mathrm{color}}\min\!\left(1,A_p(E_0)+A_p(E_1)\right).
\end{aligned}'''
    gaussian_flat = r'''\Delta_p=\left\lVert R_p(U\cup E_0)-R_p(U\cup E_1)\right\rVert_\infty,\qquad
\Delta_p\le C_{\mathrm{color}}\min\!\left(1,A_p(E_0)+A_p(E_1)\right).'''
    source = source.replace(criticality_aligned, criticality_flat)
    source = source.replace(gaussian_aligned, gaussian_flat)

    # Avoid font-dependent math-dagger fallback boxes in Word tables/notes.
    source = source.replace(r'$^\dagger$', '†')
    source = source.replace(r'\setminus', '∖')

    source = source.replace(r'\begin{minipage}{\columnwidth}', '').replace(r'\end{minipage}', '')
    word_source = build / 'word-source.tex'
    word_source.write_text(source)
    raw = build / 'manuscript-raw.docx'
    run(['pandoc', str(word_source), '--from=latex', '--bibliography=references.bib',
         '--citeproc', '--standalone', '--number-sections', '-o', str(raw)],
        ROOT, build / 'pandoc.txt')
    doc = Document(raw)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    style_map = {style.name: style for style in doc.styles}
    for style_name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4']:
        style = style_map.get(style_name)
        if style is None:
            continue
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
    for shape in doc.inline_shapes:
        if shape.width > Inches(7.1):
            ratio = Inches(7.1) / shape.width
            shape.height = int(shape.height * ratio)
            shape.width = Inches(7.1)
    widths = [[1.0, 3.8, 2.3], [1.5, .65, .8, .65, 1.0, 2.5],
              [3.1, 2.0, 2.0], [5.1, 2.0], [4.6, 2.5],
              [1.85, .6, .65, 1.0, 3.0]]
    assert len(doc.tables) == 6, f'Expected six editable tables, got {len(doc.tables)}'
    assert len(doc.inline_shapes) == 3, 'Word export must include all three figures'
    for table, sizes in zip(doc.tables, widths):
        table.autofit = False
        if 'Table Grid' in style_map:
            table.style = style_map['Table Grid']
        elif 'Table' in style_map:
            table.style = style_map['Table']
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'bottom', 'insideH']:
            border = OxmlElement('w:' + edge)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'B0B0B0')]:
                border.set(qn('w:' + key), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
        for col, width in zip(table.columns, sizes):
            col.width = Inches(width)
        for ri, row in enumerate(table.rows):
            no_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(no_split)
            if ri == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell, width in zip(row.cells, sizes):
                cell.width = Inches(width)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(2)
                    para.paragraph_format.space_before = Pt(2)
                    for run_ in para.runs:
                        run_.font.size = Pt(9)
                        if ri == 0:
                            run_.bold = True
    caption_counts = {'Image Caption': 0, 'Table Caption': 0}
    for para in doc.paragraphs:
        if para.style.name in caption_counts:
            kind = para.style.name
            caption_counts[kind] += 1
            prefix = 'Figure' if kind == 'Image Caption' else 'Table'
            run_ = OxmlElement('w:r')
            text_ = OxmlElement('w:t')
            text_.set(qn('xml:space'), 'preserve')
            text_.text = f'{prefix} {caption_counts[kind]}. '
            run_.append(text_)
            para._p.insert(1 if para._p.pPr is not None else 0, run_)
            para.paragraph_format.keep_with_next = kind == 'Table Caption'
        if para.style.name == 'Captioned Figure':
            para.paragraph_format.keep_with_next = True
    assert caption_counts == {'Image Caption': 3, 'Table Caption': 6}

    # Keep equation math editable while rendering numbers cleanly across Word/LibreOffice.
    equation_paragraphs = [p for p in doc.paragraphs if 'oMathPara' in p._p.xml]
    assert len(equation_paragraphs) == equation_count == 16
    for number, para in enumerate(equation_paragraphs, 1):
        eq_table = doc.add_table(rows=1, cols=2)
        eq_table.autofit = False
        eq_table.columns[0].width = Inches(6.35)
        eq_table.columns[1].width = Inches(0.45)
        eq_borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement('w:' + edge)
            border.set(qn('w:val'), 'nil')
            eq_borders.append(border)
        eq_table._tbl.tblPr.append(eq_borders)
        left = eq_table.cell(0, 0).paragraphs[0]
        right = eq_table.cell(0, 1).paragraphs[0]
        math_nodes = [child for child in list(para._p) if child.tag.endswith('oMathPara')]
        assert len(math_nodes) == 1
        para._p.remove(math_nodes[0])
        left._p.append(math_nodes[0])
        left.alignment = WD_ALIGN_PARAGRAPH.CENTER
        right.text = f'({number})'
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        para._p.addnext(eq_table._tbl)
        para._p.getparent().remove(para._p)

    assert not any('$$' in p.text for p in doc.paragraphs), 'Unconverted TeX equation'
    assert 'Persistent captured-world systems maintain more than' in '\n'.join(p.text for p in doc.paragraphs)
    output = ROOT / 'MAVEB_manuscript.docx'
    doc.save(output)
    with zipfile.ZipFile(output) as z:
        assert any(hashlib.sha256(z.read(n)).hexdigest() == FIGURE2_SHA256
                   for n in z.namelist() if n.startswith('word/media/'))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--build-dir', type=Path, required=True)
    parser.add_argument('--pdf-only', action='store_true')
    args = parser.parse_args()
    build = args.build_dir.resolve()
    build.mkdir(parents=True, exist_ok=True)
    image = ROOT / 'figures/figure2_exact.png'
    assert hashlib.sha256(image.read_bytes()).hexdigest() == FIGURE2_SHA256, 'Figure 2 bytes changed'
    compile_tex('main.tex', build, bibliography=True)
    compile_tex('supplement.tex', build)
    shutil.copyfile(build / 'main.pdf', ROOT / 'MAVEB_manuscript.pdf')
    shutil.copyfile(build / 'supplement.pdf', ROOT / 'MAVEB_supplement.pdf')
    if not args.pdf_only:
        export_word(build)
    print('Manuscript build and structural checks passed. Visually review before packaging.')


if __name__ == '__main__':
    main()
